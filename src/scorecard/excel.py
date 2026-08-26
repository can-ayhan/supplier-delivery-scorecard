"""Build the Excel workbook for ad hoc questions.

Three steps:

1. openpyxl writes the static sheets: an About sheet, the ranked scorecard
   and a Data sheet (one row per PO line from mart.v_delivery_lines) set up
   as an Excel table.
2. Excel itself adds the pivot tables, slicers and the trend chart on top of
   that table. openpyxl cannot create pivot tables, so this step drives Excel
   through COM and needs Windows with Excel installed.
3. The file is rewritten without the things that change on every run or
   that are specific to the machine (save time, user name), so rebuilding
   the workbook from unchanged data does not produce a different file.
"""

from __future__ import annotations

import datetime as dt
import re
import zipfile
from decimal import Decimal
from pathlib import Path

import psycopg
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

DATA_TABLE = "DeliveryLines"
NAVY = "1F3A5F"

# Flags are 0/1 so a pivot can sum them. The last two columns let a pivot
# reproduce the average lead time of the report (due lines only) as a ratio
# of two sums, which is the only kind of calculation a calculated field does.
DATA_SQL = """
    select
        supplier_name               as "Supplier",
        country_name                as "Country",
        region                      as "Region",
        material_group_name         as "Material Group",
        material_text               as "Material",
        plant_name                  as "Plant",
        po_number                   as "PO Number",
        po_item                     as "PO Item",
        order_date                  as "Order Date",
        due_date                    as "Due Date",
        due_year                    as "Due Year",
        due_quarter                 as "Due Quarter",
        due_month                   as "Due Month",
        delivery_date               as "Delivery Date",
        delivery_status             as "Delivery Status",
        uom                         as "UoM",
        ordered_qty                 as "Ordered Qty",
        received_qty                as "Received Qty",
        order_value_eur             as "Order Value EUR",
        planned_lead_time_days      as "Planned Lead Time Days",
        lead_time_days              as "Lead Time Days",
        days_late                   as "Days Late",
        due_lines                   as "Due Lines",
        on_time_lines               as "On-Time Lines",
        in_full_lines               as "In-Full Lines",
        otif_lines                  as "OTIF Lines",
        case when due_lines = 1 and lead_time_days is not null then 1 else 0 end
                                    as "Delivered Due Lines",
        case when due_lines = 1 then coalesce(lead_time_days, 0) else 0 end
                                    as "Lead Time Days (due)"
    from mart.v_delivery_lines
    order by due_date, po_number, po_item
"""

SCORECARD_SQL = """
    select
        supplier_rank       as "Rank",
        supplier_name       as "Supplier",
        country_name        as "Country",
        region              as "Region",
        due_lines           as "Due Lines",
        otif_pct            as "OTIF %",
        on_time_pct         as "On-Time %",
        in_full_pct         as "In-Full %",
        late_lines          as "Late Lines",
        avg_lead_time_days  as "Avg Lead Time (days)",
        avg_days_late       as "Avg Days Late",
        lots_inspected      as "Lots Inspected",
        defect_rate         as "Defect Rate %",
        delivery_score      as "Delivery Score",
        lead_time_score     as "Lead Time Score",
        quality_score       as "Quality Score",
        supplier_score      as "Supplier Score"
    from mart.v_supplier_scorecard
    order by supplier_rank, supplier_name
"""

FORMATS = {
    "OTIF %": "0.0%", "On-Time %": "0.0%", "In-Full %": "0.0%", "Defect Rate %": "0.00%",
    "Avg Lead Time (days)": "0.0", "Avg Days Late": "0.0", "Delivery Score": "0.0",
    "Lead Time Score": "0.0", "Quality Score": "0.0", "Supplier Score": "0.0",
    "Ordered Qty": "#,##0.###", "Received Qty": "#,##0.###", "Order Value EUR": "#,##0.00",
    "Order Date": "yyyy-mm-dd", "Due Date": "yyyy-mm-dd", "Delivery Date": "yyyy-mm-dd",
}

DEFINITIONS = [
    ("Due line", "PO line with a due date on or before the as-of date. All delivery KPIs use due lines only."),
    ("In full", "Net receipts (101 minus 102) reach the ordered quantity minus the in-full tolerance."),
    ("On time", "The receipt that completed the line was posted on or before the due date."),
    ("OTIF %", "Due lines that are both on time and in full, divided by due lines."),
    ("Lead time", "Order date to the day the line was complete, in calendar days."),
    ("Defect rate", "Blocked plus returned quantity divided by inspected quantity, for lots with a usage decision."),
    ("Supplier score", "Weighted delivery, lead time and quality scores (0-100). Weights and caps are listed below."),
]


def _plain(value):
    return float(value) if isinstance(value, Decimal) else value


def _query(conn: psycopg.Connection, sql: str) -> tuple[list[str], list[tuple]]:
    with conn.cursor() as cur:
        cur.execute(sql)
        return [d.name for d in cur.description], [tuple(_plain(v) for v in row) for row in cur.fetchall()]


def _write_sheet(ws, header: list[str], rows: list[tuple]) -> None:
    ws.append(header)
    for row in rows:
        ws.append(row)
    for index, name in enumerate(header, start=1):
        letter = get_column_letter(index)
        longest = max([len(name)] + [len(str(r[index - 1])) for r in rows[:200] if r[index - 1] is not None])
        ws.column_dimensions[letter].width = min(max(10, longest + 2), 44)
        if name in FORMATS:
            for cell in ws[letter][1:]:
                cell.number_format = FORMATS[name]
    ws.freeze_panes = "A2"


def write_base(conn: psycopg.Connection, path: Path) -> dt.date:
    """Write About, Scorecard and Data. Returns the as-of date of the data."""
    params_header, params_rows = _query(conn, "select * from mart.kpi_params")
    params = dict(zip(params_header, params_rows[0]))
    as_of = params["as_of_date"]

    wb = Workbook()

    about = wb.active
    about.title = "About"
    about["A1"] = "Supplier Delivery Scorecard"
    about["A1"].font = Font(size=16, bold=True, color=NAVY)
    about["A2"] = f"Data as of {as_of:%d %b %Y}. Built from the mart schema by `python -m scorecard excel`."
    about["A4"] = "Sheets"
    about["A4"].font = Font(bold=True)
    sheets = [
        ("Scorecard", "Ranked suppliers, same numbers as the Power BI scorecard page."),
        ("By supplier", "Pivot: delivery KPIs per supplier. Use the slicers to cut by region, material group, plant or year."),
        ("By month", "Pivot and chart: OTIF trend by due month."),
        ("By material group", "Pivot: material group, then supplier."),
        ("Data", "One row per PO line. Source of all three pivots; refresh them after rebuilding."),
    ]
    row = 5
    for name, text in sheets:
        about.cell(row=row, column=1, value=name)
        about.cell(row=row, column=2, value=text)
        row += 1
    row += 1
    about.cell(row=row, column=1, value="Definitions").font = Font(bold=True)
    for name, text in DEFINITIONS:
        row += 1
        about.cell(row=row, column=1, value=name)
        about.cell(row=row, column=2, value=text)
    row += 2
    about.cell(row=row, column=1, value="Parameters (mart.kpi_params)").font = Font(bold=True)
    for name, value in params.items():
        if name == "only_row":
            continue
        row += 1
        about.cell(row=row, column=1, value=name)
        cell = about.cell(row=row, column=2, value=value)
        cell.alignment = Alignment(horizontal="left")
    about.column_dimensions["A"].width = 30
    about.column_dimensions["B"].width = 110

    scorecard = wb.create_sheet("Scorecard")
    header, rows = _query(conn, SCORECARD_SQL)
    _write_sheet(scorecard, header, rows)
    for cell in scorecard[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    last = len(rows) + 1
    letters = {name: get_column_letter(i) for i, name in enumerate(header, start=1)}
    scorecard.conditional_formatting.add(
        f"{letters['Supplier Score']}2:{letters['Supplier Score']}{last}",
        ColorScaleRule(start_type="min", start_color="F4B6A6", mid_type="percentile", mid_value=50,
                       mid_color="FFF2CC", end_type="max", end_color="A9D6B8"),
    )
    scorecard.conditional_formatting.add(
        f"{letters['OTIF %']}2:{letters['OTIF %']}{last}",
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="2E86AB"),
    )
    scorecard.conditional_formatting.add(
        f"{letters['Defect Rate %']}2:{letters['Defect Rate %']}{last}",
        DataBarRule(start_type="num", start_value=0, end_type="max", color="C8553D"),
    )

    data = wb.create_sheet("Data")
    header, rows = _query(conn, DATA_SQL)
    _write_sheet(data, header, rows)
    table = Table(displayName=DATA_TABLE, ref=f"A1:{get_column_letter(len(header))}{len(rows) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    data.add_table(table)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return as_of


# ---------------------------------------------------------------------------
# pivot tables (Excel via COM)
# ---------------------------------------------------------------------------

XL_DATABASE, XL_ROW_FIELD, XL_SUM, XL_DESCENDING, XL_LINE_MARKERS, XL_COLUMN = 1, 1, -4157, 2, 65, 51

# calculated field -> (formula, caption, number format)
CALCULATED = {
    "OTIF Rate": ("='OTIF Lines'/'Due Lines'", "OTIF %", "0.0%"),
    "On-Time Rate": ("='On-Time Lines'/'Due Lines'", "On-Time %", "0.0%"),
    "In-Full Rate": ("='In-Full Lines'/'Due Lines'", "In-Full %", "0.0%"),
    "Lead Time": ("='Lead Time Days (due)'/'Delivered Due Lines'", "Avg Lead Time (days)", "0.0"),
}
SLICERS = ["Region", "Material Group", "Plant", "Due Year"]


def _pivot(cache, sheet, name: str, rows: list[str], values: list[str], define_calculated: bool = False):
    pivot = cache.CreatePivotTable(TableDestination=sheet.Range("A4"), TableName=name)
    if define_calculated:                # calculated fields belong to the cache, define them once
        for field, (formula, _, _) in CALCULATED.items():
            pivot.CalculatedFields().Add(field, formula, True)
    pivot.DisplayErrorString = True      # groups without due lines divide by zero
    pivot.ErrorString = ""
    for position, field in enumerate(rows, start=1):
        pivot.PivotFields(field).Orientation = XL_ROW_FIELD
        pivot.PivotFields(field).Position = position
    for value in values:
        if value in CALCULATED:
            _, caption, number_format = CALCULATED[value]
            data_field = pivot.AddDataField(pivot.PivotFields(value), caption, XL_SUM)
            data_field.NumberFormat = number_format
        else:
            data_field = pivot.AddDataField(pivot.PivotFields(value), f"{value} ", XL_SUM)
            data_field.NumberFormat = "#,##0"
    pivot.RowAxisLayout(1)               # tabular
    pivot.TableStyle2 = "PivotStyleMedium2"
    return pivot


def add_pivots(path: Path) -> None:
    """Open the workbook in Excel and add pivot tables, slicers and the trend chart."""
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:  # pragma: no cover - depends on the platform
        raise RuntimeError("the pivot step needs Windows, Excel and pywin32 (pip install .[excel])") from exc

    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(str(path.resolve()))
        cache = wb.PivotCaches().Create(SourceType=XL_DATABASE, SourceData=DATA_TABLE)

        def sheet(name: str, after: str):
            ws = wb.Worksheets.Add(None, wb.Worksheets(after))   # (Before, After)
            ws.Name = name
            return ws

        by_supplier = sheet("By supplier", "Scorecard")
        by_month = sheet("By month", "By supplier")
        by_group = sheet("By material group", "By month")

        pivots = [
            _pivot(cache, by_supplier, "ptSupplier", ["Supplier"],
                   ["Due Lines", "OTIF Rate", "On-Time Rate", "In-Full Rate", "Lead Time", "Order Value EUR"],
                   define_calculated=True),
            _pivot(cache, by_month, "ptMonth", ["Due Month"],
                   ["Due Lines", "OTIF Rate", "On-Time Rate", "In-Full Rate"]),
            _pivot(cache, by_group, "ptMaterialGroup", ["Material Group", "Supplier"],
                   ["Due Lines", "OTIF Rate", "Lead Time"]),
        ]
        pivots[0].PivotFields("Supplier").AutoSort(XL_DESCENDING, "OTIF %")

        for ws, heading in ((by_supplier, "Delivery KPIs by supplier"), (by_month, "OTIF trend by due month"),
                            (by_group, "Material group, then supplier")):
            ws.Range("A1").Value = heading
            ws.Range("A1").Font.Size = 14
            ws.Range("A1").Font.Bold = True
            ws.Range("A2").Value = "Due lines only. Slicers on the 'By supplier' sheet filter all three pivots."
            ws.Columns("A:H").AutoFit()

        left = by_supplier.Range("I4").Left
        for position, field in enumerate(SLICERS):
            slicer_cache = wb.SlicerCaches.Add2(pivots[0], field)
            for other in pivots[1:]:
                slicer_cache.PivotTables.AddPivotTable(other)
            slicer = slicer_cache.Slicers.Add(by_supplier)
            # set the layout as properties: the optional arguments of Add are ignored through late binding
            slicer.Caption = field
            slicer.Width, slicer.Height = 150, 140
            slicer.Left = left + (position % 2) * 160
            slicer.Top = by_supplier.Range("I4").Top + (position // 2) * 150

        chart_shape = by_month.Shapes.AddChart2(-1, XL_LINE_MARKERS, by_month.Range("G4").Left,
                                                by_month.Range("G4").Top, 560, 300)
        chart = chart_shape.Chart
        chart.SetSourceData(pivots[1].TableRange1)
        chart.HasTitle = True
        chart.ChartTitle.Text = "OTIF, on-time and in-full by due month"
        volume = chart.FullSeriesCollection(1)   # Due Lines: columns on their own axis,
        volume.ChartType = XL_COLUMN             # otherwise the counts flatten the percentages
        volume.AxisGroup = 2

        wb.Worksheets("About").Activate()
        wb.Save()
        wb.Close(SaveChanges=False)
    finally:
        excel.Quit()
        pythoncom.CoUninitialize()


# ---------------------------------------------------------------------------
# stable output
# ---------------------------------------------------------------------------

def _excel_serial(day: dt.date) -> int:
    return (day - dt.date(1899, 12, 30)).days


def normalise(path: Path, as_of: dt.date) -> None:
    """Replace save times, user names and local paths with values derived from the data.

    Excel stamps the file with when it was saved and by whom (document
    properties, every pivot cache) and with the folder it was saved in
    (workbook.xml). Without this the workbook differs on every build and
    carries the name and directory layout of whoever built it.
    """
    stamp = f"{as_of:%Y-%m-%d}T00:00:00Z"
    author = "Supplier Delivery Scorecard"
    replacements = {
        "docProps/core.xml": [
            (r"<dc:creator>.*?</dc:creator>", f"<dc:creator>{author}</dc:creator>"),
            (r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", f"<cp:lastModifiedBy>{author}</cp:lastModifiedBy>"),
            (r"(<dcterms:created[^>]*>).*?(</dcterms:created>)", rf"\g<1>{stamp}\g<2>"),
            (r"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)", rf"\g<1>{stamp}\g<2>"),
        ],
        "xl/workbook.xml": [
            # <x15ac:absPath url="C:\Users\...\excel\"/>, wrapped in an AlternateContent block
            (r"<mc:AlternateContent[^>]*>\s*<mc:Choice[^>]*>\s*<x15ac:absPath[^>]*/>\s*</mc:Choice>\s*</mc:AlternateContent>", ""),
        ],
    }
    cache_rules = [
        (r'refreshedBy="[^"]*"', f'refreshedBy="{author}"'),
        (r'refreshedDate="[^"]*"', f'refreshedDate="{_excel_serial(as_of)}"'),
    ]
    fixed_time = (as_of.year, as_of.month, as_of.day, 0, 0, 0)

    source = zipfile.ZipFile(path)
    entries = [(info, source.read(info.filename)) for info in source.infolist()]
    source.close()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for info, content in entries:
            rules = replacements.get(info.filename, [])
            if re.fullmatch(r"xl/pivotCache/pivotCacheDefinition\d+\.xml", info.filename):
                rules = cache_rules
            if rules:
                text = content.decode("utf-8")
                for pattern, replacement in rules:
                    text = re.sub(pattern, replacement, text, flags=re.S)
                content = text.encode("utf-8")
            entry = zipfile.ZipInfo(info.filename, date_time=fixed_time)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = info.external_attr
            target.writestr(entry, content)


def build_workbook(conn: psycopg.Connection, path: Path, pivots: bool = True) -> Path:
    as_of = write_base(conn, path)
    if pivots:
        add_pivots(path)
    normalise(path, as_of)
    return path
