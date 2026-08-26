from __future__ import annotations

import datetime as dt
import zipfile

import pytest

pytest.importorskip("openpyxl")

from openpyxl import load_workbook  # noqa: E402

from scorecard.excel import DATA_TABLE, normalise, write_base  # noqa: E402


def test_base_workbook_matches_the_marts(built, tmp_path):
    path = tmp_path / "scorecard.xlsx"

    as_of = write_base(built, path)

    wb = load_workbook(path)
    assert wb.sheetnames == ["About", "Scorecard", "Data"]
    assert as_of == dt.date(2026, 7, 31)

    data = wb["Data"]
    assert DATA_TABLE in data.tables
    lines = built.execute("select count(*) from mart.fact_po_line").fetchone()[0]
    assert data.max_row == lines + 1

    header = [cell.value for cell in data[1]]
    due = header.index("Due Lines")
    otif = header.index("OTIF Lines")
    rows = list(data.iter_rows(min_row=2, values_only=True))
    expected = built.execute(
        "select count(*) filter (where is_due), count(*) filter (where is_due and is_otif) from mart.fact_po_line"
    ).fetchone()
    assert (sum(r[due] for r in rows), sum(r[otif] for r in rows)) == expected

    scorecard = wb["Scorecard"]
    assert scorecard.max_row == 23                      # 22 suppliers and the header
    assert scorecard["A2"].value == 1                   # best rank first


def test_normalise_removes_user_and_save_time(tmp_path):
    path = tmp_path / "book.xlsx"
    core = (
        "<cp:coreProperties><dc:creator>someone</dc:creator><cp:lastModifiedBy>Some One</cp:lastModifiedBy>"
        '<dcterms:created xsi:type="dcterms:W3CDTF">2030-01-02T03:04:05Z</dcterms:created>'
        '<dcterms:modified xsi:type="dcterms:W3CDTF">2030-01-02T03:04:05Z</dcterms:modified></cp:coreProperties>'
    )
    cache = '<pivotCacheDefinition refreshedBy="Some One" refreshedDate="47485.127" refreshedVersion="8"/>'
    workbook = (
        '<workbook xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><fileVersion appName="xl"/>'
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:Choice Requires="x15"><x15ac:absPath url="C:\\Users\\someone\\reports\\" '
        'xmlns:x15ac="http://schemas.microsoft.com/office/spreadsheetml/2010/11/ac"/></mc:Choice>'
        "</mc:AlternateContent><sheets/></workbook>"
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("docProps/core.xml", core)
        archive.writestr("xl/pivotCache/pivotCacheDefinition1.xml", cache)
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/styles.xml", "<styleSheet/>")

    normalise(path, dt.date(2026, 7, 31))

    with zipfile.ZipFile(path) as archive:
        core = archive.read("docProps/core.xml").decode()
        cache = archive.read("xl/pivotCache/pivotCacheDefinition1.xml").decode()
        workbook = archive.read("xl/workbook.xml").decode()
        assert archive.read("xl/styles.xml") == b"<styleSheet/>"       # everything else is left alone
        assert {info.date_time for info in archive.infolist()} == {(2026, 7, 31, 0, 0, 0)}
    assert "Some One" not in core + cache and "someone" not in core + workbook
    assert "absPath" not in workbook and '<fileVersion appName="xl"/><sheets/>' in workbook
    assert core.count("2026-07-31T00:00:00Z") == 2
    assert 'refreshedDate="46234"' in cache and 'refreshedVersion="8"' in cache
